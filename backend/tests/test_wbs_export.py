"""WBS 的两个导出出口：Excel 表 与 调试框图。

盯住的是几条「看不出报错」的事：导出的表里分组行和叶子行混着，字号分不分层、
排除了几条有没有写出来、分组行的状态是不是被随手填了一个值；框图那边则是
版面只有一份（页面与 Excel 必须是同一个 spec 算出来的）。

**不在模块顶层 import 应用模块**（见 CLAUDE.md）：`conftest.py` 的 client 夹具靠
os.chdir 把 sqlite 指到临时目录，收集阶段的顶层 import 会赶在 chdir 之前把引擎
连上仓库里的 backend/app.db。
"""
import io

import pytest


@pytest.fixture(scope="module")
def special_id(client, admin_headers):
    r = client.post("/api/specials", json={"name": "导出测试专项"}, headers=admin_headers)
    assert r.status_code in (200, 201), r.text
    return r.json()["id"]


@pytest.fixture()
def plan(client, admin_headers, special_id):
    """一份两层的 WBS：两个阶段，各挂两个叶子，其中一个标「已变更」。"""
    r = client.post("/api/wbs/plans",
                    json={"kind": "special", "special_id": special_id, "name": "导出用 WBS"},
                    headers=admin_headers)
    assert r.status_code == 200, r.text
    pid = r.json()["id"]
    for stage in ("调试准备", "功能验证"):
        d = client.post(f"/api/wbs/plans/{pid}/items", json={"name": stage},
                        headers=admin_headers).json()
        sid = [x for x in d["items"] if x["name"] == stage][0]["id"]
        for j, (nm, st, days) in enumerate((("子任务甲", "进行中", 3.0),
                                            ("子任务乙", "已变更", 5.0))):
            client.post(f"/api/wbs/plans/{pid}/items",
                        json={"parent_id": sid, "name": f"{stage}-{nm}",
                              "status": st, "man_days": days, "progress_pct": 40 * (j + 1)},
                        headers=admin_headers)
    return pid


# ─── 调试框图的版面 ─────────────────────────────────────────────────────────
def test_diagram_columns_are_the_top_level_rows(client, admin_headers, plan):
    """第 1 层＝一列。列数和阶段数对不上的话，图上就少了一整段活。"""
    spec = client.get(f"/api/wbs/plans/{plan}/diagram", headers=admin_headers).json()
    assert spec["stage_count"] == 2
    assert spec["box_count"] == 6            # 2 个阶段 + 4 个叶子
    # 相邻阶段之间一根箭头：n 个阶段 n-1 根
    assert len(spec["arrows"]) == 1


def test_diagram_font_size_steps_down_with_depth(client, admin_headers, plan):
    """各级字号必须真的分档——这正是「一目了然」要的那件事。

    盯的是「第 2 层比第 1 层小」，不是具体几号：阶梯改了这条仍然成立，
    而两级一样大（谁不小心把 wbs_level_font 改成常数）立刻红。
    """
    spec = client.get(f"/api/wbs/plans/{plan}/diagram", headers=admin_headers).json()
    by_depth = {}
    for b in spec["boxes"]:
        by_depth.setdefault(b["depth"], set()).add(b["font_px"])
    assert len(by_depth[1]) == 1 and len(by_depth[2]) == 1
    assert next(iter(by_depth[1])) > next(iter(by_depth[2]))


def test_diagram_keeps_changed_rows_and_greys_them(client, admin_headers, plan):
    """「已变更」的行照画、上灰底。剔掉的话，图上看不出这件事为什么没做。"""
    import brand

    spec = client.get(f"/api/wbs/plans/{plan}/diagram", headers=admin_headers).json()
    grey = [b for b in spec["boxes"] if b["fill"] == "#" + brand.STATUS_FILLS["已变更"]]
    assert len(grey) == 2
    assert {"label": "已变更", "fill": "#" + brand.STATUS_FILLS["已变更"]} in spec["legend"]


def test_diagram_legend_only_lists_status_words_actually_on_the_chart(client, admin_headers, plan):
    """图例只列图上真出现过的状态：全档铺出来会让人以为自己漏看了绿的那几个。"""
    spec = client.get(f"/api/wbs/plans/{plan}/diagram", headers=admin_headers).json()
    labels = {x["label"] for x in spec["legend"]}
    assert labels == {"进行中", "已变更"}


def test_diagram_reports_what_it_could_not_draw():
    """画不下的条数要如实报出来（同 nodeClipped / unassigned / match_rate）。"""
    import wbs_diagram

    rows = []
    for i in range(1, 60):
        rows.append({"code": str(i), "name": f"阶段{i}", "depth": 1, "is_leaf": False,
                     "leaf_count": 3, "days": 3.0, "pct": 0, "status": "", "owner": ""})
        for j in range(1, 6):
            rows.append({"code": f"{i}.{j}", "name": f"任务{j}", "depth": 2, "is_leaf": True,
                         "leaf_count": 1, "days": 1.0, "pct": 0, "status": "未开始", "owner": ""})
    spec = wbs_diagram.build_diagram(rows, "大树")
    assert spec["box_count"] < len(rows)
    assert spec["skipped"] == len(rows) - spec["box_count"]
    # 截断按**整段阶段**截：从一个阶段中间切开的话，图上那一列看着就是"这个阶段
    # 就这么多活"，而它其实还有一半没画
    codes = {b["code"].split(".")[0] for b in spec["boxes"]}
    for c in codes:
        assert sum(1 for b in spec["boxes"] if b["code"].split(".")[0] == c) == 6


def test_diagram_wraps_long_names_inside_the_box():
    """名字长到一行装不下要折行，不能戳出盒子——盒子看着还规规矩矩，字却在外面。"""
    import wbs_diagram

    long_name = "一个特别特别长的调试任务名字长到一行根本放不下必须折行才行"
    spec = wbs_diagram.build_diagram(
        [{"code": "1", "name": long_name, "depth": 1, "is_leaf": True,
          "leaf_count": 1, "days": 0, "pct": 0, "status": "", "owner": ""}], "折行")
    box = spec["boxes"][0]
    assert len(box["lines"]) > 1
    inner = box["w"] - 2 * wbs_diagram.BOX_PAD_X
    for ln in box["lines"]:
        assert wbs_diagram._text_w(ln, box["font_px"]) <= inner + 0.01


# ─── Excel 导出 ────────────────────────────────────────────────────────────
def _load(client, headers, plan_id):
    import openpyxl

    r = client.get(f"/api/wbs/plans/{plan_id}/export.xlsx", headers=headers)
    assert r.status_code == 200, r.text
    return openpyxl.load_workbook(io.BytesIO(r.content))


def test_export_has_every_row_including_the_changed_ones(client, admin_headers, plan):
    """导出**不剔**「已变更」的行：那是交付记录，也是"为什么没做"的唯一线索。"""
    wb = _load(client, admin_headers, plan)
    ws = wb["WBS"]
    names = [ws.cell(r, 2).value for r in range(2, 8)]
    assert len([n for n in names if n]) == 6
    assert any("子任务乙" in (n or "") for n in names)


def test_export_font_size_steps_down_with_depth(client, admin_headers, plan):
    """表里的字号也按层级分档，且与页面同一份阶梯（enums.wbs_level_font）。"""
    import enums

    wb = _load(client, admin_headers, plan)
    ws = wb["WBS"]
    sizes = {}
    for r in range(2, 8):
        code = str(ws.cell(r, 1).value or "")
        depth = code.count(".") + 1
        sizes.setdefault(depth, set()).add(ws.cell(r, 2).font.size)
    assert sizes[1] == {enums.wbs_level_font(1)["pt"]}
    assert sizes[2] == {enums.wbs_level_font(2)["pt"]}
    assert next(iter(sizes[1])) > next(iter(sizes[2]))


def test_export_leaves_group_status_blank(client, admin_headers, plan):
    """分组行的状态**留空**。随手填个「进行中」会被当成有人标过——
    而状态是叶子填的，分组没有状态（同「未指定领域」那行的采集问题单留空）。"""
    wb = _load(client, admin_headers, plan)
    ws = wb["WBS"]
    for r in range(2, 8):
        is_group = "分组" in str(ws.cell(r, 3).value or "")
        if is_group:
            assert not (ws.cell(r, 10).value or "")
        else:
            assert ws.cell(r, 10).value


def test_export_says_how_many_rows_it_excluded(client, admin_headers, plan):
    """只筛不报的数字比没有更糟：排除了几条、多少人天要写在表尾。"""
    wb = _load(client, admin_headers, plan)
    ws = wb["WBS"]
    tail = "\n".join(str(ws.cell(r, 1).value or "") for r in range(8, ws.max_row + 1))
    assert "已排除 2 条" in tail
    assert "10" in tail            # 两条各 5 人天
    assert "仍然留在表里" in tail


def test_export_carries_the_diagram_sheet(client, admin_headers, plan):
    """第 2 页是框图，且说明与版面是同一份 spec 算出来的。"""
    wb = _load(client, admin_headers, plan)
    assert wb.sheetnames == ["WBS", "调试框图"]
    ws2 = wb["调试框图"]
    head = "\n".join(str(ws2.cell(r, 1).value or "") for r in range(1, 5))
    assert "调试框图" in head
    assert "2 个阶段" in head and "6 个方框" in head


def test_empty_plan_exports_without_blaming_the_font(client, admin_headers, special_id):
    """空 WBS 与「这台机器没中文字体」是两回事，混成一句会让人白装一遍字体。"""
    r = client.post("/api/wbs/plans",
                    json={"kind": "special", "special_id": special_id, "name": "空 WBS"},
                    headers=admin_headers)
    pid = r.json()["id"]
    wb = _load(client, admin_headers, pid)
    for sheet in ("调试框图",):
        text = "\n".join(str(wb[sheet].cell(i, 1).value or "") for i in range(1, 8))
        assert "还没有任何工作包" in text, sheet
        assert "字体" not in text, sheet


# ─── 框图里的计划日期与延期 ───────────────────────────────────────────────
@pytest.fixture()
def dated(client, admin_headers, special_id):
    """一份带计划日期的 WBS：一条已过期没做完、一条按期、一条只填完成日、一条什么都没填。"""
    r = client.post("/api/wbs/plans",
                    json={"kind": "special", "special_id": special_id, "name": "带日期的 WBS"},
                    headers=admin_headers)
    pid = r.json()["id"]
    d = client.post(f"/api/wbs/plans/{pid}/items", json={"name": "阶段甲"},
                    headers=admin_headers).json()
    sid = [x for x in d["items"] if x["name"] == "阶段甲"][0]["id"]
    rows = [
        ("拖了的活", "进行中", "2026-01-05", "2026-01-20", 40),
        ("按期的活", "进行中", "2099-12-01", "2099-12-20", 10),
        ("已完成但过了期的活", "已完成", "2026-01-05", "2026-01-20", 100),
        ("已变更且过了期的活", "已变更", "2026-01-05", "2026-01-20", 0),
        ("只填完成日的活", "未开始", None, "2099-12-25", 0),
        ("什么都没填的活", "未开始", None, None, 0),
    ]
    for nm, st, s0, s1, pct in rows:
        client.post(f"/api/wbs/plans/{pid}/items",
                    json={"parent_id": sid, "name": nm, "status": st, "man_days": 2,
                          "progress_pct": pct, "planned_start": s0, "planned_end": s1},
                    headers=admin_headers)
    return pid


def _boxes(client, headers, pid, **params):
    d = client.get(f"/api/wbs/plans/{pid}/diagram", headers=headers, params=params).json()
    return d, {b["lines"][0].split(" ", 1)[1]: b for b in d["boxes"] if b["is_leaf"]}


def test_every_box_carries_its_planned_date(client, admin_headers, dated):
    """方框第三行是计划日期——这正是它和一份事务清单的区别。"""
    _, by = _boxes(client, admin_headers, dated)
    assert by["按期的活"]["date"] == "计划 12-01 → 12-20"
    assert by["只填完成日的活"]["date"] == "计划 完成 12-25"


def test_box_marks_only_real_overdue_rows(client, admin_headers, dated):
    """延期＝已过计划完成日 且 状态不是「已完成」 且 不是「已变更 / 不涉及」。

    这三条哪一条放宽了，「延期 N 个」里就会混进一批根本不该追的活，
    而那个数看着还挺合理，没人会去核。
    """
    d, by = _boxes(client, admin_headers, dated)
    late = {k for k, b in by.items() if b["overdue"]}
    assert late == {"拖了的活"}
    assert "已延期" in by["拖了的活"]["date"]
    assert d["overdue"] == 1          # **只数叶子**，父行跟着红但不另计


def test_overdue_changes_the_border_not_the_fill(client, admin_headers, dated):
    """延期改**边框**不改底色：底色表达的是状态，拿它表示延期就得在
    「进行中」和「已延期」里二选一，而那正是要同时看到的两件事。"""
    import brand

    _, by = _boxes(client, admin_headers, dated)
    b = by["拖了的活"]
    assert b["stroke_w"] > 1 and b["stroke"] != b["fill"]
    assert b["fill"] == "#" + brand.STATUS_FILLS["进行中"]     # 底色仍是状态色


def test_box_says_so_when_the_due_date_is_missing(client, admin_headers, dated):
    """没填计划完成日的如实写出来，不留空：留空会被读成"没有交期要求"，
    而那批行正是最该被追着去补的（同 overdue_unknown）。"""
    d, by = _boxes(client, admin_headers, dated)
    assert by["什么都没填的活"]["date"] == "未填计划完成日"
    assert d["undated"] == 1
    assert "没填计划完成日" in " ".join(d["note_lines"])


def test_uncounted_rows_get_no_date_line_and_no_progress_bar(client, admin_headers, dated):
    """「已变更 / 不涉及」的行不写日期、不画完成度条：那条本轮就不做了，
    交期无从谈起；画一条空槽出来看着像"一点没做"，而它其实是不做了。"""
    _, by = _boxes(client, admin_headers, dated)
    b = by["已变更且过了期的活"]
    assert b["date"] == "" and b["bar"] is None and b["overdue"] is False


def test_chart_and_table_agree_on_what_is_overdue(client, admin_headers, dated):
    """图上标红的那几条，必须正好是表格里写着「已过计划完成日」的那几条。

    两处各写一份判定的表现是**表格里写着延期、图上不红**，而两边单独看都对。
    """
    detail = client.get(f"/api/wbs/plans/{dated}", headers=admin_headers).json()
    flagged = {it["name"] for it in detail["items"]
               if any("已过计划完成日" in x for x in it["issues"])}
    _, by = _boxes(client, admin_headers, dated)
    assert {k for k, b in by.items() if b["overdue"]} == flagged


def test_parent_box_is_red_only_when_a_leaf_under_it_is(client, admin_headers, dated):
    """父行的红从叶子推上来，不是拿汇总的完成日再判一次——后者会把
    「子任务都按期、只是整段跨到了今天之后」也标成延期。"""
    d, _ = _boxes(client, admin_headers, dated)
    parent = [b for b in d["boxes"] if not b["is_leaf"]][0]
    assert parent["overdue"] is True          # 底下有「拖了的活」
    assert parent["date"].startswith("汇总 ")   # 父行的日期明写是汇总值


# ─── 导出到第几层 ──────────────────────────────────────────────────────────
def test_export_depth_keeps_the_totals_whole(client, admin_headers, plan):
    """只导出到第 1 层时，**合计一个数都不能变**：砍的是"列出来的层级"，
    不是"算进去的活"。拿裁剪后的行去算合计，导出到第 1 层会得到一个只数了
    2 行的人天，而那个数看着还挺合理。
    """
    import openpyxl

    def totals(md):
        r = client.get(f"/api/wbs/plans/{plan}/export.xlsx", headers=admin_headers,
                       params={"max_depth": md} if md else {})
        ws = openpyxl.load_workbook(io.BytesIO(r.content))["WBS"]
        tail = "\n".join(str(ws.cell(i, 1).value or "") for i in range(2, ws.max_row + 1))
        body = sum(1 for i in range(2, ws.max_row + 1)
                   if str(ws.cell(i, 3).value or "") in ("工作包",)
                   or "分组" in str(ws.cell(i, 3).value or ""))
        return [ln for ln in tail.split("\n") if ln.startswith("合计：")][0], body, tail

    full, full_rows, _ = totals(None)
    clipped, clipped_rows, tail = totals(1)
    assert full == clipped                       # 合计一字不差
    assert clipped_rows < full_rows              # 但列出来的行少了
    assert "只导出到第 1 层" in tail
    assert "仍然把它们算在内" in tail            # 只筛不报比没有更糟


def test_diagram_depth_clip_reports_what_it_folded(client, admin_headers, plan):
    """图也吃同一次裁剪，且折叠了几行要报出来——只筛不报的表现是
    「这个阶段怎么只剩一个方框了」，而没人说得清少的是哪些。"""
    d = client.get(f"/api/wbs/plans/{plan}/diagram", headers=admin_headers,
                   params={"max_depth": 1}).json()
    assert d["box_count"] == 2                   # 两个阶段，子任务都折叠了
    assert d["folded"] == 4 and d["max_depth"] == 1
