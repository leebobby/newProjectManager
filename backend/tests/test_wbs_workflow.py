"""WBS 日常填报、富文本详情及依赖显示模式。"""
import importlib.util
import io
from pathlib import Path
from uuid import uuid4

import pytest


@pytest.fixture
def workflow(client, admin_headers):
    def post(path, data):
        result = client.post('/api' + path, headers=admin_headers, json=data)
        assert result.status_code == 200, result.text
        return result.json()
    special = post('/specials', {'name': '日常填报回归'})
    plan = post('/wbs/plans', {'name': '日常填报', 'kind': 'special', 'special_id': special['id']})
    for name in ('准备', '验证', '交付'):
        detail = post(f"/wbs/plans/{plan['id']}/items", {'name': name, 'man_days': 1,
                       'planned_start': '2026-10-01', 'planned_end': '2026-10-20'})
    return detail


def test_owner_selection_fills_group_and_respects_explicit_group(client, admin_headers, workflow):
    code = uuid4().hex[:12]
    department = client.post('/api/resource-groups', headers=admin_headers,
                             json={'code': code + 'd', 'name': '填报部门', 'kind': 'dept'})
    assert department.status_code == 200, department.text
    response = client.post('/api/resource-groups', headers=admin_headers,
                           json={'code': code, 'name': '填报组', 'kind': 'pl', 'parent_id': department.json()['id']})
    assert response.status_code == 200, response.text
    group = response.json()
    user = client.post('/api/users', headers=admin_headers,
                       json={'username': code, 'full_name': '责任人', 'can_login': False,
                             'group_id': group['id']}).json()
    item = workflow['items'][0]
    result = client.put(f"/api/wbs/items/{item['id']}", headers=admin_headers,
                        json={'owner_user_id': user['id'], 'version': item['version']})
    assert result.status_code == 200, result.text
    saved = result.json()['items'][0]
    assert (saved['owner'], saved['group_id'], saved['owner_group']) == ('责任人', group['id'], '填报组')
    result = client.put(f"/api/wbs/items/{item['id']}", headers=admin_headers,
                        json={'owner_user_id': user['id'], 'group_id': None, 'version': saved['version']})
    assert result.json()['items'][0]['group_id'] is None
    no_group = client.post('/api/users', headers=admin_headers,
                           json={'username': code + 'n', 'full_name': '未分组人员', 'can_login': False}).json()
    client.put(f"/api/wbs/items/{item['id']}", headers=admin_headers, json={'group_id': group['id']})
    result = client.put(f"/api/wbs/items/{item['id']}", headers=admin_headers,
                        json={'owner_user_id': no_group['id']})
    assert result.json()['items'][0]['group_id'] == group['id']


def test_progress_100_completes_leaf_and_updates_rollup(client, admin_headers, workflow):
    item = workflow['items'][0]
    result = client.put(f"/api/wbs/items/{item['id']}", headers=admin_headers,
                        json={'progress_pct': 100, 'version': item['version']})
    assert result.status_code == 200
    saved = result.json()['items'][0]
    assert saved['status'] == '已完成' and saved['roll_pct'] == 100
    assert result.json()['progress_pct'] == 33
    assert not any('延期' in issue for issue in saved['issues'])
    created = client.post(f"/api/wbs/plans/{workflow['id']}/items", headers=admin_headers,
                          json={'name': '已做完', 'progress_pct': 100, 'status': '进行中'})
    assert created.json()['items'][-1]['status'] == '已完成'


def test_details_round_trip_multiple_predecessors_rich_text_and_risks(client, admin_headers, workflow):
    first, second, target = workflow['items']
    risk = {'kind': '风险', 'description': '待验证接口', 'owner': '张明', 'progress': '已安排验证',
            'raised_on': '2026-10-01', 'resolved_on': '2026-10-10'}
    payload = {'version': target['version'], 'current_progress': '<p><b>联调完成</b></p><img src=x onerror=alert(1)>',
               'issues_risks': [risk], 'predecessor_ids': [first['id'], second['id']],
               'deliverable': '旧交付物', 'remark': '旧备注'}
    result = client.put(f"/api/wbs/items/{target['id']}", headers=admin_headers, json=payload)
    assert result.status_code == 200, result.text
    saved = result.json()['items'][-1]
    assert '<b>联调完成</b>' in saved['current_progress']
    assert 'onerror' not in saved['current_progress']
    assert saved['issues_risks'] == [risk]
    assert [p['id'] for p in saved['predecessors']] == [first['id'], second['id']]
    read = client.get(f"/api/wbs/plans/{workflow['id']}", headers=admin_headers).json()['items'][-1]
    assert read['current_progress'] == saved['current_progress'] and read['issues_risks'] == [risk]
    exported = client.get(f"/api/wbs/plans/{workflow['id']}/export.xlsx", headers=admin_headers)
    import openpyxl
    sheet = openpyxl.load_workbook(io.BytesIO(exported.content)).worksheets[0]
    assert sheet.cell(4, 16).value == '联调完成'
    assert '待验证接口' in sheet.cell(4, 17).value
    assert '2026-10-10' in sheet.cell(4, 17).value
    clear = client.put(f"/api/wbs/items/{target['id']}", headers=admin_headers,
                       json={'version': saved['version'], 'issues_risks': [], 'current_progress': ''})
    cleared = clear.json()['items'][-1]
    assert cleared['issues_risks'] == [] and cleared['current_progress'] == ''
    assert cleared['deliverable'] == '旧交付物' and cleared['remark'] == '旧备注'


def test_risk_dates_reject_inverted_range(client, admin_headers, workflow):
    item = workflow['items'][0]
    result = client.put(f"/api/wbs/items/{item['id']}", headers=admin_headers,
                        json={'issues_risks': [{'raised_on': '2026-10-20', 'resolved_on': '2026-10-01'}]})
    assert result.status_code == 422
    assert client.get(f"/api/wbs/plans/{workflow['id']}/timeline", headers=admin_headers,
                      params={'dependency_mode': 'invalid'}).status_code == 422


@pytest.mark.parametrize('mode', ['arrows', 'lines', 'hidden'])
def test_dependency_modes_preserve_relations_and_match_excel(client, admin_headers, workflow, mode, monkeypatch):
    import wbs_timeline
    first, second, target = workflow['items']
    client.put(f"/api/wbs/items/{target['id']}", headers=admin_headers,
               json={'predecessor_ids': [first['id'], second['id']]})
    url = f"/api/wbs/plans/{workflow['id']}"
    spec = client.get(url + '/timeline', params={'dependency_mode': mode}, headers=admin_headers).json()
    assert spec['dependency_mode'] == mode
    assert len(spec['arrows']) == (0 if mode == 'hidden' else 2)
    assert spec['arrowheads'] == (mode == 'arrows')
    captured = []
    monkeypatch.setattr(wbs_timeline, 'render_png', lambda layout: captured.append(layout))
    response = client.get(url + '/export.xlsx', params={'dependency_mode': mode}, headers=admin_headers)
    assert response.status_code == 200
    assert captured[0]['arrows'] == spec['arrows'] and captured[0]['arrowheads'] == spec['arrowheads']
    detail = client.get(url, headers=admin_headers).json()
    assert len(detail['items'][-1]['predecessors']) == 2


def test_existing_database_migration_preserves_legacy_data(tmp_path):
    import sqlalchemy as sa
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    script = Path(__file__).resolve().parents[1] / 'alembic/versions/0020_wbs_progress_risks.py'
    spec = importlib.util.spec_from_file_location('wbs_progress_migration', script)
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    engine = sa.create_engine('sqlite:///' + str(tmp_path / 'legacy.db'))
    with engine.begin() as connection:
        connection.exec_driver_sql('CREATE TABLE wbs_items (id INTEGER PRIMARY KEY, remark TEXT)')
        connection.exec_driver_sql("INSERT INTO wbs_items VALUES (1, '原有风险备注')")
        with Operations.context(MigrationContext.configure(connection)):
            migration.upgrade()
            migration.upgrade()
        row = connection.exec_driver_sql('SELECT remark, current_progress, issues_risks_json FROM wbs_items').one()
        assert tuple(row) == ('原有风险备注', '', '[]')
